from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import re

def create_ppt_with_template(txt_file, template_file, output_file):
    prs = Presentation(template_file)
    
    with open(txt_file, 'r', encoding='utf-8') as f:
        content = f.read()
        
    slides_text = content.split('\n\n')
    
    for slide_content in slides_text:
        if not slide_content.strip(): 
            continue
            
        slide = prs.slides.add_slide(prs.slide_layouts[0]) 
        
        # 기본 찌꺼기 상자 지우기
        for shape in slide.shapes:
            sp = shape.element
            sp.getparent().remove(sp)
            
        # 1. 텍스트 상자 크기 조정 (화면 좌우 꽉 채우기)
        left = Inches(0)         # 왼쪽 여백을 0으로 만듭니다.
        top = Inches(0.5)        # 위쪽 여백은 그대로 유지합니다.
        width = prs.slide_width  # 너비를 슬라이드 전체 너비와 똑같이 맞춥니다.
        height = prs.slide_height - Inches(1)
        
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP 
        
        lines = slide_content.strip().split('\n')
        
        for i, line in enumerate(lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
                
            p.text = line.strip()
            p.alignment = PP_ALIGN.CENTER
            
            # 한글 포함 여부 확인
            is_korean = bool(re.search(r'[가-힣]', line))
            
            # 2. 한글/영어에 따른 폰트 및 크기 세팅
            font_size = 45 if is_korean else 40
            font_name = 'Malgun Gothic' if is_korean else 'Arial'
            
            for run in p.runs:
                run.font.size = Pt(font_size)
                # run.font.bold = True # (선택) 맑은 고딕 Regular를 위해 굵게(Bold) 설정을 지웠습니다. 필요시 복구하세요.
                run.font.name = font_name
                run.font.color.rgb = RGBColor(255, 255, 255) # 흰색 글씨
                
    prs.save(output_file)
    print("폰트와 너비가 완벽하게 세팅된 PPT 완성! 폴더를 확인해보세요.")

# 스크립트 실행
create_ppt_with_template('lyrics.txt', 'template.pptx', 'final_praise.pptx')