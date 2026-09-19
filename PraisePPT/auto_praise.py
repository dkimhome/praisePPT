from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import re

def make_ultimate_praise_ppt(kor_file, eng_file, template_file, output_file):
    # ==========================================
    # 1. 가사 믹스 단계 (2줄씩 묶기)
    # ==========================================
    with open(kor_file, 'r', encoding='utf-8') as f:
        kor_lines = [line.strip() for line in f.readlines() if line.strip()]
        
    with open(eng_file, 'r', encoding='utf-8') as f:
        eng_lines = [line.strip() for line in f.readlines() if line.strip()]
        
    mixed_content = []
    
    # 2줄씩 건너뛰면서(step=2) 가사를 가져와 묶어줍니다.
    for i in range(0, max(len(kor_lines), len(eng_lines)), 2):
        # 한글 2줄, 영어 2줄 가져오기
        kor_part = "\n".join(kor_lines[i:i+2])
        eng_part = "\n".join(eng_lines[i:i+2])
        
        # 슬라이드 하나에 들어갈 내용 합치기
        slide_text = f"{kor_part}\n{eng_part}".strip()
        
        if slide_text:
            mixed_content.append(slide_text)
            
    # 나중에 눈으로 확인할 수 있게 중간 결과물(lyrics_mixed.txt)도 하나 저장해 둡니다.
    with open('lyrics_mixed.txt', 'w', encoding='utf-8') as f:
        f.write('\n\n'.join(mixed_content))

    # ==========================================
    # 2. PPT 디자인 단계 
    # ==========================================
    prs = Presentation(template_file)
    
    for slide_content in mixed_content:
        slide = prs.slides.add_slide(prs.slide_layouts[0]) 
        
        for shape in slide.shapes:
            sp = shape.element
            sp.getparent().remove(sp)
            
        # 텍스트 상자 크기 (좌우 꽉 차게)
        left = Inches(0)         
        top = Inches(0.5)        
        width = prs.slide_width  
        height = prs.slide_height - Inches(1)
        
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP 
        
        lines = slide_content.split('\n')
        
        for i, line in enumerate(lines):
            if not line.strip():
                continue
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
                
            p.text = line.strip()
            p.alignment = PP_ALIGN.CENTER
            
            # 한글/영어 구분 및 폰트 적용
            is_korean = bool(re.search(r'[가-힣]', line))
            font_size = 45 if is_korean else 40
            font_name = 'Malgun Gothic' if is_korean else 'Arial'
            # 2. 색상 조정: 영어를 더 화사하고 밝은 오렌지색(255, 220, 60)으로 변경
            font_color = RGBColor(255, 255, 255) if is_korean else RGBColor(255, 220, 60)
            
            for run in p.runs:
                run.font.size = Pt(font_size)
                run.font.name = font_name
                run.font.color.rgb = font_color
                
    prs.save(output_file)
    print("🎉 가사 믹스부터 PPT 완성까지 단 1초 만에 완료되었습니다!")

# 스크립트 실행 명령 (파일 이름들을 넣어줍니다)
make_ultimate_praise_ppt('korean.txt', 'english.txt', 'template.pptx', 'final_praise.pptx')